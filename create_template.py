"""Create a compact, ATS-safe single-page DOCX resume template."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt


def _configure_style(document: Document, style_name: str, size: float) -> None:
    """Configure a compact, readable paragraph style."""
    style = document.styles[style_name]
    style.font.name = "Aptos"
    style.font.size = Pt(size)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0


def _tabs(paragraph, position: float = 7.05) -> None:
    """Add a right tab stop for date and location alignment."""
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(position), WD_TAB_ALIGNMENT.RIGHT)


def _paragraph(document: Document, text: str = "", style: str = "Normal"):
    """Create a tightly spaced paragraph with consistent layout."""
    paragraph = document.add_paragraph(text, style=style)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    return paragraph


def _heading(document: Document, title: str) -> None:
    """Add a compact section heading with visual separation."""
    paragraph = _paragraph(document)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(9)
    run.font.all_caps = True


def _bullet(document: Document, text: str) -> None:
    """Add a dense ATS-readable bullet paragraph."""
    paragraph = _paragraph(document, text, "List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.16)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)


def build_template() -> Path:
    """Build and save the compact DOCX template used by document export."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.38)
    section.bottom_margin = Inches(0.38)
    section.left_margin = Inches(0.52)
    section.right_margin = Inches(0.52)
    _configure_style(document, "Normal", 8.8)
    _configure_style(document, "List Bullet", 8.4)

    name = _paragraph(document)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(1)
    run = name.add_run("{{ name }}")
    run.bold = True
    run.font.size = Pt(16)

    for text in ["{{ email }} | {{ phone }} | {{ location }}", "{{ linkedin }} | {{ github }} | {{ portfolio }}"]:
        line = _paragraph(document, text)
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    divider = _paragraph(document, "_" * 114)
    divider.paragraph_format.space_after = Pt(1)

    _heading(document, "Professional Summary")
    _paragraph(document, "{{ summary }}")

    _heading(document, "Education")
    _paragraph(document, "{% for edu in education %}")
    row = _paragraph(document)
    _tabs(row)
    row.add_run("{{ edu.institution }}").bold = True
    row.add_run("\t{{ edu.location }}")
    row = _paragraph(document)
    _tabs(row)
    row.add_run("{{ edu.degree }} | CGPA: {{ edu.cgpa }}")
    row.add_run("\t{{ edu.years }}")
    _paragraph(document, "{% endfor %}")

    _heading(document, "Technical Skills")
    skills = _paragraph(document)
    skills.add_run("Languages: ").bold = True
    skills.add_run("{{ languages }}  |  ")
    skills.add_run("Tools: ").bold = True
    skills.add_run("{{ tools }}  |  ")
    skills.add_run("Databases: ").bold = True
    skills.add_run("{{ databases }}")

    _heading(document, "Internship Experience")
    _paragraph(document, "{% for exp in experience %}")
    row = _paragraph(document)
    _tabs(row)
    row.add_run("{{ exp.role }}").bold = True
    row.add_run("\t{{ exp.dates }}")
    _paragraph(document, "{{ exp.company }} | {{ exp.location }}")
    _paragraph(document, "{% for bullet in exp.bullets %}")
    _bullet(document, "{{ bullet }}")
    _paragraph(document, "{% endfor %}")
    _paragraph(document, "{% endfor %}")

    _heading(document, "Projects")
    _paragraph(document, "{% for project in projects %}")
    row = _paragraph(document)
    _tabs(row)
    row.add_run("{{ project.name }} | {{ project.stack }}").bold = True
    row.add_run("\t{{ project.date }}")
    _paragraph(document, "{% for bullet in project.bullets %}")
    _bullet(document, "{{ bullet }}")
    _paragraph(document, "{% endfor %}")
    _paragraph(document, "{% endfor %}")

    _heading(document, "Achievements & Certifications")
    _paragraph(document, "{% for achievement in achievements %}")
    _bullet(document, "{{ achievement }}")
    _paragraph(document, "{% endfor %}")

    target = Path(__file__).resolve().parent / "templates" / "resume_template.docx"
    target.parent.mkdir(exist_ok=True)
    document.save(target)
    return target


if __name__ == "__main__":
    print(build_template())
